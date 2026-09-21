"""
Reliability and Human-in-the-Loop Escalation Gate.

Safeguards document verification by identifying when the system cannot confidently
decide on a shipment comparison, escalating to an operations human reviewer rather
than guessing or failing silently.

Escalation Reasons:
- missing_attachment: Document comparison requested, but BL attachment is missing or dropped.
- wrong_doc_type: Attached document is a Commercial Invoice, Packing List, or COO instead of a BL.
- unreadable: Scanned image-only PDF without OCR layer, empty file (0-byte), or corrupted bytes.
- missing_value: Key comparison field in SI is left blank or filled with placeholder tokens.
"""

import os
from pathlib import Path
from typing import Dict, Any, Tuple, Optional, List

try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None

try:
    import docx
except ImportError:
    docx = None

try:
    import openpyxl
except ImportError:
    openpyxl = None

BLANK_TOKENS = {
    "???",
    "_______",
    "______",
    "____",
    "TBA",
    "TBC",
    "N/A",
    "____MT",
    "",
}

KEY_FIELD_LABELS = [
    "SHIPPER",
    "CONSIGNEE",
    "NOTIFY",
    "LOADING",
    "LOAD PORT",
    "POL",
    "DISCHARGE",
    "POD",
    "CONTAINER",
    "GROSS WEIGHT",
    "GROSS WT",
]


class ReliabilityGate:
    """Evaluates whether an email/attachment pair can be safely automated or needs human review."""

    def __init__(self, data_root: str = "data_v2"):
        self.data_root = Path(data_root)

    def evaluate(self, email_record: Dict[str, Any]) -> Tuple[bool, Optional[str], Dict[str, Any]]:
        """
        Evaluates the email record and its attachments for reliability issues.

        Args:
            email_record: Dict containing email metadata and attachments list.

        Returns:
            (needs_review: bool, review_reason: Optional[str], details: Dict[str, Any])
        """
        attachments: List[str] = email_record.get("attachments", [])
        body = (email_record.get("body") or "").lower()
        details: Dict[str, Any] = {"attachments": attachments}

        # 1. Missing Attachment Check
        if len(attachments) == 1:
            details["reason_explanation"] = "Only one attachment provided (expected SI and draft BL pair)."
            return True, "missing_attachment", details

        if len(attachments) == 0:
            if any(term in body for term in ["dropped", "missing", "forgot", "left off"]):
                details["reason_explanation"] = "Email indicates attachments were intended but none were attached."
                return True, "missing_attachment", details
            # Normal pending document request (e.g. "please assist to send draft BL")
            return False, None, details

        # 2. Unreadable / Corrupted File Check
        doc_texts: Dict[str, str] = {}
        for att_rel in attachments:
            att_path = self.data_root / att_rel
            if not att_path.exists():
                details["reason_explanation"] = f"Attachment file not found: {att_rel}"
                return True, "unreadable", details

            try:
                size = att_path.stat().st_size
                if size == 0:
                    details["reason_explanation"] = f"Attachment file is empty (0 bytes): {att_rel}"
                    return True, "unreadable", details
            except Exception as e:
                details["reason_explanation"] = f"Error reading file status: {e}"
                return True, "unreadable", details

            ext = att_path.suffix.lower()
            text_extracted = ""

            if ext == ".txt":
                try:
                    with open(att_path, "r", encoding="utf-8", errors="replace") as f:
                        text_extracted = f.read()
                except Exception as e:
                    details["reason_explanation"] = f"Failed to decode text file {att_rel}: {e}"
                    return True, "unreadable", details

            elif ext == ".pdf":
                if fitz is None:
                    details["reason_explanation"] = "PyMuPDF (fitz) is required to inspect PDF documents."
                    return True, "unreadable", details
                try:
                    doc = fitz.open(str(att_path))
                    text_extracted = "".join(page.get_text() for page in doc).strip()
                    if len(text_extracted) == 0:
                        # Image-only scanned PDF without OCR layer
                        details["reason_explanation"] = (
                            f"Scanned image PDF without machine-readable text layer: {att_rel}"
                        )
                        return True, "unreadable", details
                except Exception as e:
                    details["reason_explanation"] = f"Failed to parse PDF document {att_rel}: {e}"
                    return True, "unreadable", details

            elif ext == ".docx":
                if docx is None:
                    details["reason_explanation"] = "python-docx is required to inspect Word documents."
                    return True, "unreadable", details
                try:
                    d = docx.Document(str(att_path))
                    parts = [p.text for p in d.paragraphs]
                    for tbl in d.tables:
                        for row in tbl.rows:
                            parts.extend(c.text for c in row.cells)
                    text_extracted = "\n".join(parts)
                except Exception as e:
                    details["reason_explanation"] = f"Failed to parse Word document {att_rel}: {e}"
                    return True, "unreadable", details

            elif ext == ".xlsx":
                if openpyxl is None:
                    details["reason_explanation"] = "openpyxl is required to inspect Excel documents."
                    return True, "unreadable", details
                try:
                    wb = openpyxl.load_workbook(str(att_path))
                    ws = wb.active
                    parts = []
                    for row in ws.iter_rows(values_only=True):
                        for cell in row:
                            if cell is not None:
                                parts.append(str(cell))
                    text_extracted = " ".join(parts)
                except Exception as e:
                    details["reason_explanation"] = f"Failed to parse Excel document {att_rel}: {e}"
                    return True, "unreadable", details

            doc_texts[att_rel] = text_extracted

        # 3. Wrong Document Type Check
        bl_candidates = [a for a in attachments if "_BL." in a]
        target_bl = bl_candidates[0] if bl_candidates else attachments[1]
        bl_content = doc_texts.get(target_bl, "").upper()

        if "COMMERCIAL INVOICE" in bl_content:
            details["reason_explanation"] = (
                f"Attachment {target_bl} is a Commercial Invoice, not a draft Bill of Lading."
            )
            return True, "wrong_doc_type", details

        if "PACKING LIST" in bl_content:
            details["reason_explanation"] = (
                f"Attachment {target_bl} is a Packing List, not a draft Bill of Lading."
            )
            return True, "wrong_doc_type", details

        if "CERTIFICATE OF ORIGIN" in bl_content:
            details["reason_explanation"] = (
                f"Attachment {target_bl} is a Certificate of Origin, not a draft Bill of Lading."
            )
            return True, "wrong_doc_type", details

        # 4. Missing / Blank Values in Required Fields
        si_candidates = [a for a in attachments if "_SI." in a]
        target_si = si_candidates[0] if si_candidates else attachments[0]
        si_content = doc_texts.get(target_si, "")

        for line in si_content.splitlines():
            if ":" in line:
                parts = line.split(":", 1)
                lbl = parts[0].strip().upper()
                val = parts[1].strip()
                if any(k in lbl for k in KEY_FIELD_LABELS):
                    clean_val = val.replace(" ", "")
                    if val in BLANK_TOKENS or clean_val in BLANK_TOKENS:
                        details["reason_explanation"] = (
                            f"Required field '{parts[0].strip()}' in {target_si} has blank or unpopulated token: '{val}'"
                        )
                        return True, "missing_value", details

        return False, None, details
