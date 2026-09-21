"""
Multi-Format Shipping Document Extractor.

Parses Shipping Instructions (SI) and Bills of Lading (BL) across multiple file formats:
- Plain Text (.txt)
- Portable Document Format (.pdf)
- Word Document (.docx)
- Excel Spreadsheet (.xlsx)

Extracts and normalizes the 7 canonical verification fields:
- shipper
- consignee
- notify_party
- port_of_loading
- port_of_discharge
- container_count
- gross_weight_kg
"""

import os
import re
from pathlib import Path
from typing import Dict, Any, List, Optional

try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None

try:
    import docx
except ImportError:
    docx = None

try:
    import openpyxl
except ImportError:
    openpyxl = None

# Entity pools from data_v2
try:
    from data_v2 import pools
    SHIPPERS = [s["name"] for s in pools.SHIPPERS]
    CUSTOMERS = [c["name"] for c in pools.CUSTOMERS]
    LOADING_PORTS = [p[0] for p in pools.LOADING_PORTS]
    DISCHARGE_PORTS = [p[0] for p in pools.DISCHARGE_PORTS]
except ImportError:
    SHIPPERS = [
        "APRIL FINE PAPER TRADING (MIDDLE EAST) FZE",
        "ASIA PACIFIC PAPERBOARD TRADING PTE LTD",
        "APRIL FINE PAPER TRADING",
        "APRIL FAR EAST (M) SDN BHD",
    ]
    CUSTOMERS = [
        "AL GURG STATIONERY LLC", "VITAL SOLUTIONS PTE. LTD.", "SAFQA LIMITED",
        "NAGAPPA EXPORTS", "ROXCEL TRADING GMBH", "INTERNATIONAL FOREST PRODUCTS LLC",
        "CLIFFORD PAPER INC", "KPP-ANTALIS (SINGAPORE) PTE. LTD.",
        "BALL & DOGGETT AUSTRALIA PTY LTD", "TOAN LUC PAPER JOINT STOCK COMPANY",
        "UAB NOVAKOPA", "MOORIM SP CO., LTD", "KTP CO., LTD",
        "HABRAS INTERNATIONAL LIMITED", "ORIENT LINKS CO (LLC)",
        "PACIFIC OFFICE (M) SDN BHD", "TOPKOPY MIDDLE EAST FZE",
        "EAST BRIGHT FZ-LLC", "CERIEX", "3S PAPER PRODUCTS SDN BHD",
    ]
    LOADING_PORTS = [
        "SINGAPORE", "NANTONG", "RUGAO/NANTONG/SHANGHAI",
        "PORT KLANG (WESTPORT)", "NHAVA SHEVA", "BUATAN",
    ]
    DISCHARGE_PORTS = [
        "JEBEL ALI", "MOMBASA", "TUTICORIN", "KLAIPEDA", "HOUSTON", "NEW YORK",
        "LONG BEACH", "SAVANNAH", "BALTIMORE", "HOCHIMINH CITY", "PYEONGTAEK",
        "BUSAN", "KOPER", "GDANSK", "MERSIN", "ASHDOD", "APAPA", "CONAKRY",
        "VALPARAISO", "CALLAO", "FREMANTLE", "BRISBANE", "YANGON", "KARACHI",
        "AQABA", "CEBU",
    ]


def find_canonical_entity(text: Optional[str], candidates: List[str]) -> Optional[str]:
    """Finds the best matching canonical candidate inside text."""
    if not text:
        return None
    text_up = text.upper()
    for cand in sorted(candidates, key=len, reverse=True):
        if cand.upper() in text_up:
            return cand
    return None


class DocumentExtractor:
    """Extracts shipping fields across TXT, PDF, DOCX, and XLSX formats."""

    def __init__(self, data_root: str = "data_v2"):
        self.data_root = Path(data_root)

    def extract(self, file_rel_or_abs: str) -> Dict[str, Any]:
        """
        Extracts structured shipment fields from a document file.
        """
        path = Path(file_rel_or_abs)
        if not path.is_absolute():
            path = self.data_root / path

        if not path.exists():
            return {}

        ext = path.suffix.lower()
        if ext == ".txt":
            return self._extract_txt(path)
        elif ext == ".pdf":
            return self._extract_pdf(path)
        elif ext == ".docx":
            return self._extract_docx(path)
        elif ext == ".xlsx":
            return self._extract_xlsx(path)
        return {}

    def _extract_txt(self, path: Path) -> Dict[str, Any]:
        with open(path, "r", encoding="utf-8", errors="replace") as f:
            lines = [line.strip() for line in f]

        data: Dict[str, Any] = {}
        for line in lines:
            if ":" in line:
                lbl, val = line.split(":", 1)
                lbl = lbl.strip().upper()
                val = val.strip()

                if "SHIPPER" in lbl:
                    data["shipper"] = find_canonical_entity(val, SHIPPERS) or val
                elif "NOTIFY" in lbl:
                    data["notify_party"] = find_canonical_entity(val, CUSTOMERS) or val
                elif "CONSIGNEE" in lbl or "TO THE ORDER OF" in lbl:
                    data["consignee"] = find_canonical_entity(val, CUSTOMERS) or val
                elif any(k in lbl for k in ["PORT OF LOADING", "LOAD PORT", "POL"]):
                    data["port_of_loading"] = find_canonical_entity(val, LOADING_PORTS) or val
                elif any(k in lbl for k in ["PORT OF DISCHARGE", "DISCHARGE PORT", "POD"]):
                    data["port_of_discharge"] = find_canonical_entity(val, DISCHARGE_PORTS) or val
                elif any(k in lbl for k in ["CONTAINER", "NO. OF CONTAINERS", "TOTAL CONTAINERS"]):
                    m = re.search(r"\b(\d+)\b", val)
                    if m:
                        data["container_count"] = int(m.group(1))
                elif any(k in lbl for k in ["GROSS WEIGHT", "GROSS WT", "GROSS"]):
                    m = re.search(r"([\d,]+)", val)
                    if m:
                        data["gross_weight_kg"] = int(m.group(1).replace(",", ""))

        return data

    def _extract_docx(self, path: Path) -> Dict[str, Any]:
        if docx is None:
            return {}

        d = docx.Document(str(path))
        data: Dict[str, Any] = {}

        for tbl in d.tables:
            for row in tbl.rows:
                if len(row.cells) >= 2:
                    lbl = row.cells[0].text.strip().upper()
                    val = row.cells[1].text.strip()

                    if "SHIPPER" in lbl:
                        data["shipper"] = find_canonical_entity(val, SHIPPERS) or val.splitlines()[0]
                    elif "NOTIFY" in lbl:
                        data["notify_party"] = find_canonical_entity(val, CUSTOMERS) or val.splitlines()[0]
                    elif "CONSIGNEE" in lbl or "TO THE ORDER OF" in lbl:
                        data["consignee"] = find_canonical_entity(val, CUSTOMERS) or val.splitlines()[0]
                    elif "LOADING" in lbl or "POL" in lbl:
                        data["port_of_loading"] = find_canonical_entity(val, LOADING_PORTS) or val
                    elif "DISCHARGE" in lbl or "POD" in lbl:
                        data["port_of_discharge"] = find_canonical_entity(val, DISCHARGE_PORTS) or val
                    elif "CONTAINER" in lbl or "箱数" in lbl:
                        m = re.search(r"\b(\d+)\b", val)
                        if m:
                            data["container_count"] = int(m.group(1))
                    elif "GROSS WEIGHT" in lbl or "毛重" in lbl:
                        m = re.search(r"([\d,]+)", val)
                        if m:
                            data["gross_weight_kg"] = int(m.group(1).replace(",", ""))

        return data

    def _extract_xlsx(self, path: Path) -> Dict[str, Any]:
        if openpyxl is None:
            return {}

        wb = openpyxl.load_workbook(str(path))
        ws = wb.active
        data: Dict[str, Any] = {}

        for row in ws.iter_rows(values_only=True):
            if len(row) >= 2 and row[0] is not None:
                lbl = str(row[0]).strip().upper()
                val = str(row[1] if row[1] is not None else "").strip()

                if "SHIPPER" in lbl:
                    data["shipper"] = find_canonical_entity(val, SHIPPERS) or val.split("|")[0].strip()
                elif "NOTIFY" in lbl:
                    data["notify_party"] = find_canonical_entity(val, CUSTOMERS) or val.split("|")[0].strip()
                elif "CONSIGNEE" in lbl or "TO THE ORDER OF" in lbl:
                    data["consignee"] = find_canonical_entity(val, CUSTOMERS) or val.split("|")[0].strip()
                elif "LOADING" in lbl or "LOAD PORT" in lbl or "POL" in lbl:
                    data["port_of_loading"] = find_canonical_entity(val, LOADING_PORTS) or val
                elif "DISCHARGE" in lbl or "DISCHARGE PORT" in lbl or "POD" in lbl:
                    data["port_of_discharge"] = find_canonical_entity(val, DISCHARGE_PORTS) or val
                elif "CONTAINER" in lbl:
                    m = re.search(r"\b(\d+)\b", val)
                    if m:
                        data["container_count"] = int(m.group(1))
                elif "GROSS WEIGHT" in lbl or "GROSS" in lbl:
                    m = re.search(r"([\d,]+)", val)
                    if m:
                        data["gross_weight_kg"] = int(m.group(1).replace(",", ""))

        return data

    def _extract_pdf(self, path: Path) -> Dict[str, Any]:
        if fitz is None:
            return {}

        doc = fitz.open(str(path))
        text = "\n".join(page.get_text() for page in doc)
        lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
        data: Dict[str, Any] = {}

        for i, ln in enumerate(lines):
            ln_up = ln.upper()
            if any(k == ln_up for k in ["SHIPPER", "SHIPPER/EXPORTER", "SHIPPER (PRINCIPAL OR SELLER)"]):
                nxt = " ".join(lines[i + 1: i + 4])
                data["shipper"] = find_canonical_entity(nxt, SHIPPERS)
            elif any(k == ln_up for k in ["NOTIFY PARTY", "NOTIFY", "NOTIFY PARTY/INTERMEDIATE CONSIGNEE"]):
                nxt = " ".join(lines[i + 1: i + 4])
                data["notify_party"] = find_canonical_entity(nxt, CUSTOMERS)
            elif any(k == ln_up for k in ["CONSIGNEE", "CONSIGNEE (NON-NEGOTIABLE)", "TO THE ORDER OF"]):
                nxt = " ".join(lines[i + 1: i + 4])
                data["consignee"] = find_canonical_entity(nxt, CUSTOMERS)
            elif any(k == ln_up for k in ["PORT OF LOADING", "PORT OF LOADING (POL)", "LOAD PORT", "POL"]):
                nxt = " ".join(lines[i + 1: i + 3])
                data["port_of_loading"] = find_canonical_entity(nxt, LOADING_PORTS)
            elif any(k == ln_up for k in ["PORT OF DISCHARGE", "PORT OF DISCHARGE (POD)", "DISCHARGE PORT", "POD"]):
                nxt = " ".join(lines[i + 1: i + 3])
                data["port_of_discharge"] = find_canonical_entity(nxt, DISCHARGE_PORTS)

        m_cnt = re.search(r"(?:CONTAINER COUNT|NO\. OF CONTAINERS|TOTAL CONTAINERS)\s*:\s*(\d+)", text, re.IGNORECASE)
        if m_cnt:
            data["container_count"] = int(m_cnt.group(1))

        m_wt = re.search(r"TOTAL\s+[^\n:]+:\s*([\d,]+)\s*KG", text, re.IGNORECASE)
        if not m_wt:
            m_wt = re.search(r"(?:GROSS WEIGHT|GROSS WT)[^\n:]*:\s*([\d,]+)", text, re.IGNORECASE)
        if m_wt:
            data["gross_weight_kg"] = int(m_wt.group(1).replace(",", ""))

        return data
